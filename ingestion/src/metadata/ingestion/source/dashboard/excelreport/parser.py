#  Copyright 2025 Collate
#  Licensed under the Collate Community License, Version 1.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#  https://github.com/open-metadata/OpenMetadata/blob/main/ingestion/LICENSE
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
"""
Parses a QT.B14.QLDL.MB06.1 "YCPT_BAOCAO" style report-request workbook
into `ParsedDashboard` objects.

The template is a governance form filled in by hand by two different
teams (business + data), so sheet names, exact row offsets and helper
rows above the real header vary a bit between requests. Sheets are
therefore located by the *content* of their header row (a set of
expected Vietnamese/English column names) rather than by sheet name or
position, which is the only thing that stays constant across requests.

Layout this parser understands, based on the sample
`QT.B14.QLDL.MB06.1.YCPT_BAOCAO_TOP_10.xlsx`:
  * "Tổng quan báo cáo" sheet  -> one row per Dashboard
  * "Template Dashboard" sheet -> free-form mockup, kept as reference text only
  * "mapping báo cáo" sheet    -> one row per report field, with the
    business description columns (meaning/calc logic/unit) and the data
    columns (source schema/table/field, mapping rule) side by side. This
    is the sheet used to build Charts (grouped by "Tên Sheet", falling back
    to "Tên mục/biểu đồ" when a report has no sheet breakdown) and to
    resolve source-table lineage.
  * "Đánh giá tác động" sheet  -> optional, kept as reference text only
"""

import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional, Tuple  # noqa: UP035

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from metadata.ingestion.source.dashboard.excelreport.models import (
    MappingRow,
    ParsedDashboard,
    ReportOverview,
)
from metadata.utils.logger import ingestion_logger

logger = ingestion_logger()

OVERVIEW_MARKERS = [
    "tên báo cáo",
    "mã báo cáo",
    "loại báo cáo",
    "link báo cáo",
    "thư mục đặt báo cáo",
]
MAPPING_MARKERS = [
    "tên chỉ tiêu",
    "loại chỉ tiêu",
    "tên mục",
    "source schema",
    "source table",
    "source field",
]
MAX_HEADER_SCAN_ROWS = 20
MIN_HEADER_SCORE = 3


def normalize(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip().lower()


def clean(value) -> Optional[str]:  # noqa: UP045
    """Strip a cell value down to a plain string, dropping empties."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _as_int(value) -> Optional[int]:  # noqa: UP045
    """Best-effort parse of an STT cell into an int, without treating bools as 0/1."""
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str) and value.strip().isdigit():
        return int(value.strip())
    return None


def _is_index_legend(text: str) -> bool:
    """Matches helper rows like '(1) | (2) | (3) ...' used to number the header."""
    return bool(re.fullmatch(r"\(\d+\)", text.strip()))


@dataclass
class HeaderMatch:
    row: int
    columns: Dict[int, str]  # noqa: UP006
    score: int


def find_header_row(
    worksheet: Worksheet,
    markers: List[str],  # noqa: UP006
    max_scan_rows: int = MAX_HEADER_SCAN_ROWS,
) -> Optional[HeaderMatch]:  # noqa: UP045
    """Scans the top of a sheet for the row that looks the most like the
    expected header, by counting how many marker phrases show up in it."""
    best: Optional[HeaderMatch] = None  # noqa: UP045
    scan_limit = min(max_scan_rows, worksheet.max_row or 0)
    for row in range(1, scan_limit + 1):
        columns = {
            col: normalize(worksheet.cell(row=row, column=col).value)
            for col in range(1, (worksheet.max_column or 0) + 1)
        }
        columns = {col: text for col, text in columns.items() if text}
        if not columns:
            continue
        score = sum(
            1 for marker in markers if any(marker in text for text in columns.values())
        )
        if best is None or score > best.score:
            best = HeaderMatch(row=row, columns=columns, score=score)
    if best is None or best.score < min(MIN_HEADER_SCORE, len(markers)):
        return None
    return best


def get_col(header: HeaderMatch, *keywords: str) -> Optional[int]:  # noqa: UP045
    """Returns the leftmost column whose header text contains all keywords."""
    lowered = [keyword.lower() for keyword in keywords]
    for col in sorted(header.columns):
        text = header.columns[col]
        if all(keyword in text for keyword in lowered):
            return col
    return None


def _row_is_blank(worksheet: Worksheet, row: int) -> bool:
    max_column = (worksheet.max_column or 0) + 1
    return all(
        worksheet.cell(row=row, column=col).value in (None, "")
        for col in range(1, max_column)
    )


def _find_row_label(
    worksheet: Worksheet, row: int, stt_col: Optional[int]
) -> Optional[str]:  # noqa: UP045
    """Used for rows that are not data rows and not header noise - typically a
    'Báo cáo XYZ' section title spanning a single cell, grouping the rows below it."""
    if stt_col:
        label = clean(worksheet.cell(row=row, column=stt_col).value)
        if label:
            return label
    for col in range(1, (worksheet.max_column or 0) + 1):
        label = clean(worksheet.cell(row=row, column=col).value)
        if label:
            return label
    return None


def iter_data_rows(
    worksheet: Worksheet,
    header: HeaderMatch,
) -> List[Tuple[int, str | None]]:  # noqa: UP006
    """Walks the rows below the header, returning (row_index, current_group_label)
    for every row that looks like real data (STT column holds a number).

    Rows in between (index-legend rows, verbose help rows, section-title rows used
    to group indicators under a named report) are consumed to update the running
    group label rather than emitted as data.
    """
    stt_col = get_col(header, "stt")
    data_rows: List[Tuple[int, str | None]] = []  # noqa: UP006
    current_label: Optional[str] = None  # noqa: UP045
    for row in range(header.row + 1, (worksheet.max_row or 0) + 1):
        if _row_is_blank(worksheet, row):
            continue
        stt_value = worksheet.cell(row=row, column=stt_col).value if stt_col else None
        if _as_int(stt_value) is not None:
            data_rows.append((row, current_label))
            continue
        stt_text = normalize(stt_value)
        if stt_text == "stt" or _is_index_legend(stt_text):
            continue
        label = _find_row_label(worksheet, row, stt_col)
        if label:
            current_label = label
    return data_rows


def parse_overview_sheet(
    worksheet: Worksheet, header: HeaderMatch
) -> List[ReportOverview]:  # noqa: UP006
    col_name = get_col(header, "tên báo cáo")
    col_code = get_col(header, "mã báo cáo")
    col_folder = get_col(header, "thư mục đặt báo cáo")
    col_link = get_col(header, "link báo cáo")
    col_type = get_col(header, "loại báo cáo")
    col_unit1 = get_col(header, "đơn vị yêu cầu cấp 1")
    col_unit2 = get_col(header, "đơn vị yêu cầu cấp 2")
    col_business_contact = get_col(header, "đầu mối nghiệp vụ")
    col_data_unit = get_col(header, "đơn vị phụ trách") or get_col(header, "kdl")
    col_data_contact = get_col(header, "đầu mối kdl")
    col_purpose = get_col(header, "ý nghĩa") or get_col(header, "mục đích")
    col_using_unit = get_col(header, "đơn vị sử dụng")
    col_using_title = get_col(header, "chức danh sử dụng")
    col_permission = get_col(header, "phân quyền dữ liệu")
    col_channel = get_col(header, "khai thác trên kênh")
    col_params = get_col(header, "tham số báo cáo")
    col_frequency = get_col(header, "tần suất")
    col_date_range = get_col(header, "khoảng thời gian")
    col_history = get_col(header, "dữ liệu lịch sử")
    col_daily_time = get_col(header, "trong ngày")
    col_sensitive = get_col(header, "nhạy cảm")

    def value(row: int, col: Optional[int]) -> Optional[str]:  # noqa: UP045
        return clean(worksheet.cell(row=row, column=col).value) if col else None

    overviews: List[ReportOverview] = []  # noqa: UP006
    for row, _ in iter_data_rows(worksheet, header):
        report_name = value(row, col_name)
        if not report_name:
            continue
        overviews.append(
            ReportOverview(
                report_name=report_name,
                report_code=value(row, col_code),
                folder=value(row, col_folder),
                link=value(row, col_link),
                report_type=value(row, col_type),
                requester_unit_l1=value(row, col_unit1),
                requester_unit_l2=value(row, col_unit2),
                business_contact=value(row, col_business_contact),
                data_unit=value(row, col_data_unit),
                data_contact=value(row, col_data_contact),
                purpose=value(row, col_purpose),
                using_unit=value(row, col_using_unit),
                using_title=value(row, col_using_title),
                data_permission=value(row, col_permission),
                channel=value(row, col_channel),
                params=value(row, col_params),
                frequency=value(row, col_frequency),
                date_range=value(row, col_date_range),
                history_required=value(row, col_history),
                daily_time=value(row, col_daily_time),
                sensitive_data=value(row, col_sensitive),
            )
        )
    return overviews


def parse_mapping_sheet(
    worksheet: Worksheet, header: HeaderMatch
) -> List[MappingRow]:  # noqa: UP006
    col_indicator = get_col(header, "tên chỉ tiêu")
    col_type = get_col(header, "loại chỉ tiêu")
    col_sheet = get_col(header, "tên sheet")
    col_section = get_col(header, "tên mục")
    col_meaning = get_col(header, "ý nghĩa")
    col_logic = get_col(header, "logic tính toán")
    col_unit = get_col(header, "đơn vị tính")
    col_source_system = get_col(header, "hệ thống dữ liệu nguồn")
    col_input_ref = get_col(header, "dẫn chiếu màn hình")
    col_src_schema = get_col(header, "source schema")
    col_src_table = get_col(header, "source table")
    col_src_field = get_col(header, "source field")
    col_mapping_rule = get_col(header, "mapping rule")
    col_rule_reconcile = get_col(header, "rule reconcile")
    col_metadata = get_col(header, "metadata")
    col_note = get_col(header, "note")

    def value(row: int, col: Optional[int]) -> Optional[str]:  # noqa: UP045
        return clean(worksheet.cell(row=row, column=col).value) if col else None

    rows: List[MappingRow] = []  # noqa: UP006
    for row, group_label in iter_data_rows(worksheet, header):
        indicator_name = value(row, col_indicator)
        if not indicator_name:
            continue
        rows.append(
            MappingRow(
                indicator_name=indicator_name,
                group_label=group_label,
                indicator_type=value(row, col_type),
                sheet_name=value(row, col_sheet),
                section_name=value(row, col_section),
                meaning=value(row, col_meaning),
                calc_logic=value(row, col_logic),
                unit=value(row, col_unit),
                source_system=value(row, col_source_system),
                input_ref=value(row, col_input_ref),
                source_schema=value(row, col_src_schema),
                source_table=value(row, col_src_table),
                source_field=value(row, col_src_field),
                mapping_rule=value(row, col_mapping_rule),
                rule_reconcile=value(row, col_rule_reconcile),
                metadata_note=value(row, col_metadata),
                note=value(row, col_note),
            )
        )
    return rows


def render_sheet_as_markdown(
    worksheet: Worksheet, max_rows: int = 200
) -> Optional[str]:  # noqa: UP045
    """Best-effort, generic dump of a free-form sheet (mockup / impact assessment)
    into a Markdown table, kept as reference documentation on the Dashboard rather
    than parsed into entities - these sheets are hand-drawn and not consistently
    structured between report requests."""
    used_rows = []
    max_row = min(worksheet.max_row or 0, max_rows)
    min_col, max_col = None, None
    for row in range(1, max_row + 1):
        values = [
            clean(worksheet.cell(row=row, column=col).value)
            for col in range(1, (worksheet.max_column or 0) + 1)
        ]
        if not any(values):
            continue
        used_rows.append((row, values))
        non_empty_cols = [i for i, v in enumerate(values, start=1) if v]
        min_col = (
            min(non_empty_cols) if min_col is None else min([min_col, *non_empty_cols])
        )
        max_col = (
            max(non_empty_cols) if max_col is None else max([max_col, *non_empty_cols])
        )
    if not used_rows or min_col is None:
        return None
    lines = []
    for _row, values in used_rows:
        cells = values[min_col - 1 : max_col]
        rendered = " | ".join((cell or "").replace("\n", " ") for cell in cells)
        lines.append(f"| {rendered} |")
    return "\n".join(lines)


@dataclass
class SheetClassification:
    overview: Optional[Worksheet] = None  # noqa: UP045
    overview_header: Optional[HeaderMatch] = None  # noqa: UP045
    mapping: Optional[Worksheet] = None  # noqa: UP045
    mapping_header: Optional[HeaderMatch] = None  # noqa: UP045
    mockup: Optional[Worksheet] = None  # noqa: UP045
    impact: Optional[Worksheet] = None  # noqa: UP045


def classify_sheets(workbook) -> SheetClassification:
    classification = SheetClassification()
    remaining: List[Worksheet] = []  # noqa: UP006
    for worksheet in workbook.worksheets:
        name = normalize(worksheet.title)
        a1 = normalize(worksheet["A1"].value)
        if "tác động" in name or "tác động" in a1:
            classification.impact = worksheet
            continue
        if name.startswith("tổng hợp") or "thông tin tổng hợp" in a1:
            continue

        overview_header = find_header_row(worksheet, OVERVIEW_MARKERS)
        mapping_header = find_header_row(worksheet, MAPPING_MARKERS)
        overview_score = overview_header.score if overview_header else 0
        mapping_score = mapping_header.score if mapping_header else 0

        if (
            overview_score
            and overview_score >= mapping_score
            and (
                classification.overview_header is None
                or overview_score > classification.overview_header.score
            )
        ):
            classification.overview = worksheet
            classification.overview_header = overview_header
        elif mapping_score and (
            classification.mapping_header is None
            or mapping_score > classification.mapping_header.score
        ):
            classification.mapping = worksheet
            classification.mapping_header = mapping_header
        else:
            remaining.append(worksheet)

    for worksheet in remaining:
        if "template" in normalize(worksheet.title) or "template" in normalize(
            worksheet["A1"].value
        ):
            classification.mockup = worksheet
            break
    if classification.mockup is None and remaining:
        classification.mockup = remaining[0]

    return classification


def parse_workbook(path: Path) -> List[ParsedDashboard]:  # noqa: UP006
    """Parses a single YCPT_BAOCAO workbook into one ParsedDashboard per row of
    its report-overview sheet, attaching the mapping rows, mockup and impact
    assessment sheets that belong to the whole workbook."""
    workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    classification = classify_sheets(workbook)

    if classification.overview is None or classification.overview_header is None:
        logger.warning(
            f"Could not find a report-overview sheet in {path}, skipping file"
        )
        return []

    overviews = parse_overview_sheet(
        classification.overview, classification.overview_header
    )
    all_mapping_rows = (
        parse_mapping_sheet(classification.mapping, classification.mapping_header)
        if classification.mapping is not None
        and classification.mapping_header is not None
        else []
    )
    mockup_markdown = (
        render_sheet_as_markdown(classification.mockup)
        if classification.mockup is not None
        else None
    )
    impact_markdown = (
        render_sheet_as_markdown(classification.impact)
        if classification.impact is not None
        else None
    )

    dashboards = [
        ParsedDashboard(
            overview=overview,
            source_file=path,
            mockup_markdown=mockup_markdown,
            impact_markdown=impact_markdown,
        )
        for overview in overviews
    ]

    _assign_mapping_rows(dashboards, all_mapping_rows)
    return dashboards


def _assign_mapping_rows(
    dashboards: List[ParsedDashboard], mapping_rows: List[MappingRow]
) -> None:  # noqa: UP006
    """Assigns each mapping row to the dashboard whose report name best matches
    the row's section/group label. With a single dashboard in the workbook
    (the common case) every row goes to it."""
    if not dashboards:
        return
    if len(dashboards) == 1:
        dashboards[0].mapping_rows = mapping_rows
        return

    for row in mapping_rows:
        label = normalize(row.group_label or row.section_name or "")
        best_dashboard = None
        best_overlap = 0
        for dashboard in dashboards:
            name = normalize(dashboard.overview.report_name)
            overlap = (
                len(set(label.split()) & set(name.split())) if label and name else 0
            )
            if overlap > best_overlap:
                best_overlap = overlap
                best_dashboard = dashboard
        (best_dashboard or dashboards[0]).mapping_rows.append(row)


def slugify(text: str) -> str:
    """ASCII-safe token for use in an entity name; original text stays in displayName."""
    normalized = unicodedata.normalize("NFKD", text)
    ascii_text = normalized.encode("ascii", "ignore").decode("ascii")
    ascii_text = re.sub(r"[^A-Za-z0-9]+", "_", ascii_text).strip("_")
    return ascii_text.lower() or "section"


_OVERVIEW_FIELD_LABELS: List[Tuple[str, str]] = [  # noqa: UP006
    ("report_type", "Loại báo cáo"),
    ("requester_unit_l1", "Đơn vị yêu cầu cấp 1"),
    ("requester_unit_l2", "Đơn vị yêu cầu cấp 2"),
    ("business_contact", "Đầu mối nghiệp vụ"),
    ("data_unit", "Đơn vị phụ trách (KDL)"),
    ("data_contact", "Đầu mối KDL"),
    ("using_unit", "Đơn vị sử dụng"),
    ("using_title", "Chức danh sử dụng"),
    ("data_permission", "Phân quyền dữ liệu"),
    ("channel", "Kênh khai thác"),
    ("params", "Tham số báo cáo"),
    ("frequency", "Tần suất chạy báo cáo"),
    ("date_range", "Khoảng thời gian xuất dữ liệu"),
    ("history_required", "Yêu cầu dữ liệu lịch sử"),
    ("daily_time", "Thời gian có báo cáo trong ngày"),
    ("sensitive_data", "Dữ liệu nhạy cảm"),
    ("folder", "Thư mục đặt báo cáo"),
]


def _render_overview_table(overview: ReportOverview) -> str:
    rows = ["| Trường | Giá trị |", "| --- | --- |"]
    for field_name, label in _OVERVIEW_FIELD_LABELS:
        value = getattr(overview, field_name)
        if value:
            rows.append(f"| {label} | {value.replace(chr(10), ' ')} |")
    return "## Thông tin báo cáo\n\n" + "\n".join(rows)


def render_dashboard_description(dashboard: ParsedDashboard) -> str:
    overview = dashboard.overview
    sections = []
    if overview.purpose:
        sections.append(f"## Ý nghĩa / Mục đích\n\n{overview.purpose}")
    sections.append(_render_overview_table(overview))
    if dashboard.mockup_markdown:
        sections.append(
            f"## Template / Mockup báo cáo (PL3.2)\n\n{dashboard.mockup_markdown}"
        )
    if dashboard.impact_markdown:
        sections.append(f"## Đánh giá tác động (PL3.5)\n\n{dashboard.impact_markdown}")
    sections.append(f"_Nguồn: `{dashboard.source_file.name}`_")
    return "\n\n".join(sections)


def render_chart_description(rows: List[MappingRow]) -> str:  # noqa: UP006
    header = (
        "| Mục/biểu đồ | Tên chỉ tiêu | Loại | Ý nghĩa | Logic tính toán "
        "| Đơn vị tính | Nguồn | Mapping rule | Note |"
    )
    lines = [header, "| --- | --- | --- | --- | --- | --- | --- | --- | --- |"]
    for row in rows:
        source = (
            ".".join(
                part
                for part in (row.source_schema, row.source_table, row.source_field)
                if part
            )
            or "-"
        )
        lines.append(
            "| {} | {} | {} | {} | {} | {} | {} | {} | {} |".format(
                row.section_name or "-",
                row.indicator_name.replace("\n", " "),
                row.indicator_type or "-",
                (row.meaning or "-").replace("\n", " "),
                (row.calc_logic or "-").replace("\n", " "),
                row.unit or "-",
                source,
                row.mapping_rule or "-",
                row.note or "-",
            )
        )
    return "\n".join(lines)
